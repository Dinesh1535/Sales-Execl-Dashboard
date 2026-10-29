"""
Unit Tests for Sales Excel Dashboard
====================================
Tests data integrity, calculations, sheet schema, and dataset synchronization.
"""

import unittest
import os
import csv
import warnings

warnings.simplefilter("ignore")
import openpyxl

EXCEL_FILE = "sales Excel Dash Board.xlsx"
CSV_FILE = os.path.join("data", "sales_master_data.csv")


class TestSalesDashboardData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.has_excel = os.path.exists(EXCEL_FILE)
        if cls.has_excel:
            cls.wb_data = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            cls.wb_formula = openpyxl.load_workbook(EXCEL_FILE, data_only=False)

    def test_workbook_sheets_exist(self):
        """Verify that all core functional sheets exist in workbook."""
        self.assertTrue(self.has_excel, "Excel file does not exist")
        required_sheets = ["About", "Salesmans", "Master Data", "Products", "Dash Board", "Pivot Tables"]
        for sheet in required_sheets:
            self.assertIn(sheet, self.wb_data.sheetnames, f"Missing required sheet: {sheet}")

    def test_master_data_row_count(self):
        """Verify the exact count of records in Master Data."""
        ws = self.wb_data["Master Data"]
        rows = list(ws.iter_rows(values_only=True))
        # 1 header row + 278 data rows = 279 rows
        self.assertEqual(len(rows), 279, "Expected 278 data rows plus 1 header row")

    def test_formula_amount_equals_price_times_units(self):
        """Verify that Amount is strictly calculated as Price * Units for all rows."""
        ws = self.wb_data["Master Data"]
        rows = list(ws.iter_rows(values_only=True))[1:]
        for idx, row in enumerate(rows, start=2):
            price, units, amount = row[4], row[5], row[6]
            self.assertIsNotNone(price, f"Row {idx}: Price is None")
            self.assertIsNotNone(units, f"Row {idx}: Units is None")
            self.assertIsNotNone(amount, f"Row {idx}: Amount is None")
            calculated = round(price * units, 2)
            actual = round(amount, 2)
            self.assertAlmostEqual(
                calculated,
                actual,
                places=2,
                msg=f"Row {idx}: Amount mismatch ({price} * {units} = {calculated} != {actual})"
            )

    def test_no_null_values_in_master_data(self):
        """Verify that there are no empty or null values in any required columns."""
        ws = self.wb_data["Master Data"]
        rows = list(ws.iter_rows(values_only=True))[1:]
        for idx, row in enumerate(rows, start=2):
            for col_idx, val in enumerate(row):
                self.assertIsNotNone(val, f"Row {idx}, Col {col_idx+1} is None")

    def test_date_range_validity(self):
        """Verify that transaction dates fall within expected 2020-2022 range."""
        ws = self.wb_data["Master Data"]
        rows = list(ws.iter_rows(values_only=True))[1:]
        for idx, row in enumerate(rows, start=2):
            date_val = row[0]
            self.assertIsNotNone(date_val)
            year = date_val.year if hasattr(date_val, "year") else int(str(date_val)[:4])
            self.assertIn(year, [2020, 2021, 2022], f"Row {idx}: Date {date_val} outside 2020-2022")

    def test_positive_pricing_and_volume(self):
        """Verify that all prices and unit quantities are positive numbers."""
        ws = self.wb_data["Master Data"]
        rows = list(ws.iter_rows(values_only=True))[1:]
        for idx, row in enumerate(rows, start=2):
            price, units = row[4], row[5]
            self.assertGreater(price, 0, f"Row {idx}: Price must be positive")
            self.assertGreater(units, 0, f"Row {idx}: Units must be positive")

    def test_csv_export_sync(self):
        """Verify that CSV export in data/ matches Master Data row count and total sales."""
        self.assertTrue(os.path.exists(CSV_FILE), f"Missing CSV file: {CSV_FILE}")
        with open(CSV_FILE, "r", encoding="utf-8") as f:
            reader = list(csv.DictReader(f))
        self.assertEqual(len(reader), 278, "CSV must have exactly 278 records")
        csv_total_sales = sum(float(r["Amount"]) for r in reader)
        self.assertAlmostEqual(csv_total_sales, 118731893.00, places=2)


if __name__ == "__main__":
    unittest.main()
