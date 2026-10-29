"""
Sales Dashboard Analytics & Audit Tool
======================================
Automated data pipeline, validation engine, and reporting utility
for the Sales Excel Dashboard project.

Author: Antigravity & Dinesh A
"""

import sys
import os
import csv
import warnings
from datetime import datetime
from collections import defaultdict

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

EXCEL_FILE = "sales Excel Dash Board.xlsx"
CSV_FILE = os.path.join("data", "sales_master_data.csv")


def load_from_excel(filepath=EXCEL_FILE):
    """Load data from Excel workbook using openpyxl."""
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if "Master Data" not in wb.sheetnames:
        raise ValueError("Sheet 'Master Data' not found in workbook.")

    ws = wb["Master Data"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    data = []
    for r in rows[1:]:
        d, sp, prod, place, price, units, amt = r
        if d is None:
            continue
        d_str = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)
        data.append({
            "Date": d_str,
            "Sales_Person": str(sp).strip() if sp else "",
            "Product": str(prod).strip() if prod else "",
            "Place": str(place).strip() if place else "",
            "Price": float(price) if price is not None else 0.0,
            "Units": int(units) if units is not None else 0,
            "Amount": float(amt) if amt is not None else 0.0,
        })
    return data


def load_from_csv(filepath=CSV_FILE):
    """Load data from CSV file."""
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"CSV file not found: {filepath}")

    data = []
    with open(filepath, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            data.append({
                "Date": row["Date"],
                "Sales_Person": row["Sales_Person"],
                "Product": row["Product"],
                "Place": row["Place"],
                "Price": float(row["Price"]),
                "Units": int(row["Units"]),
                "Amount": float(row["Amount"]),
            })
    return data


def get_dataset():
    """Load dataset, preferring Excel workbook, falling back to CSV."""
    if os.path.exists(EXCEL_FILE):
        try:
            return load_from_excel(EXCEL_FILE)
        except Exception:
            pass
    if os.path.exists(CSV_FILE):
        return load_from_csv(CSV_FILE)
    raise RuntimeError("Neither Excel workbook nor CSV dataset could be loaded.")


def validate_dataset(records):
    """Perform comprehensive data integrity checks."""
    issues = []
    for idx, r in enumerate(records, start=1):
        # 1. Price * Units == Amount check
        expected_amt = round(r["Price"] * r["Units"], 2)
        actual_amt = round(r["Amount"], 2)
        if abs(expected_amt - actual_amt) > 0.01:
            issues.append(f"Row {idx}: Amount calculation mismatch ({r['Price']} * {r['Units']} = {expected_amt} vs {actual_amt})")

        # 2. Positive units and price
        if r["Units"] <= 0:
            issues.append(f"Row {idx}: Non-positive units ({r['Units']})")
        if r["Price"] <= 0:
            issues.append(f"Row {idx}: Non-positive price ({r['Price']})")

        # 3. Non-empty required fields
        for field in ["Sales_Person", "Product", "Place", "Date"]:
            if not r[field]:
                issues.append(f"Row {idx}: Missing field '{field}'")

    return issues


def compute_metrics(records):
    """Compute high-level summary KPIs and dimension breakdowns."""
    if not records:
        return {}

    total_amount = sum(r["Amount"] for r in records)
    total_units = sum(r["Units"] for r in records)
    total_txns = len(records)
    avg_order_value = total_amount / total_txns if total_txns else 0

    dates = [r["Date"] for r in records if r["Date"]]
    date_min = min(dates) if dates else "N/A"
    date_max = max(dates) if dates else "N/A"

    # Aggregations
    by_person = defaultdict(lambda: {"amount": 0.0, "units": 0, "txns": 0})
    by_product = defaultdict(lambda: {"amount": 0.0, "units": 0, "txns": 0})
    by_place = defaultdict(lambda: {"amount": 0.0, "units": 0, "txns": 0})
    by_year = defaultdict(lambda: {"amount": 0.0, "units": 0, "txns": 0})

    for r in records:
        sp = r["Sales_Person"]
        prod = r["Product"]
        pl = r["Place"]
        yr = r["Date"][:4]
        amt = r["Amount"]
        u = r["Units"]

        for d, k in [(by_person, sp), (by_product, prod), (by_place, pl), (by_year, yr)]:
            d[k]["amount"] += amt
            d[k]["units"] += u
            d[k]["txns"] += 1

    return {
        "kpi": {
            "total_revenue": total_amount,
            "total_units": total_units,
            "total_transactions": total_txns,
            "avg_order_value": avg_order_value,
            "date_range": f"{date_min} to {date_max}",
        },
        "sales_persons": by_person,
        "products": by_product,
        "places": by_place,
        "years": by_year,
    }


def format_currency(val):
    """Format currency in Indian Rupee format or standard comma notation."""
    return f"Rs {val:,.2f}"


def print_dashboard(metrics, validation_issues):
    """Render executive terminal dashboard."""
    kpi = metrics["kpi"]
    tot_rev = kpi["total_revenue"]

    print("=" * 72)
    print("               SALES EXCEL DASHBOARD - ANALYTICS REPORT                 ")
    print("=" * 72)
    print(f"  Reporting Period     : {kpi['date_range']}")
    print(f"  Total Revenue        : {format_currency(tot_rev)}")
    print(f"  Total Units Sold     : {kpi['total_units']:,} units")
    print(f"  Total Transactions   : {kpi['total_transactions']:,}")
    print(f"  Average Order Value  : {format_currency(kpi['avg_order_value'])}")
    print("-" * 72)

    # Health check
    if not validation_issues:
        print("  Data Integrity Status: 100% HEALTHY (0 errors found across all rows)")
    else:
        print(f"  Data Integrity Status: WARNING - {len(validation_issues)} issue(s) detected")
    print("=" * 72)

    # 1. Sales Person Performance
    print("\n[+] SALES REPRESENTATIVE PERFORMANCE LEADERBOARD")
    print(f"{'Sales Person':<15} | {'Total Sales':<16} | {'Units':<8} | {'Orders':<7} | {'Share %':<8}")
    print("-" * 65)
    for sp, data in sorted(metrics["sales_persons"].items(), key=lambda x: x[1]["amount"], reverse=True):
        share = (data["amount"] / tot_rev) * 100
        print(f"{sp:<15} | {format_currency(data['amount']):<16} | {data['units']:<8} | {data['txns']:<7} | {share:>6.2f}%")

    # 2. Product Performance
    print("\n[+] PRODUCT CATEGORY PERFORMANCE")
    print(f"{'Product':<15} | {'Total Sales':<16} | {'Units':<8} | {'Orders':<7} | {'Share %':<8}")
    print("-" * 65)
    for p, data in sorted(metrics["products"].items(), key=lambda x: x[1]["amount"], reverse=True):
        share = (data["amount"] / tot_rev) * 100
        print(f"{p:<15} | {format_currency(data['amount']):<16} | {data['units']:<8} | {data['txns']:<7} | {share:>6.2f}%")

    # 3. Regional Breakdown
    print("\n[+] REGIONAL / PLACE PERFORMANCE")
    print(f"{'Region (Place)':<15} | {'Total Sales':<16} | {'Units':<8} | {'Orders':<7} | {'Share %':<8}")
    print("-" * 65)
    for pl, data in sorted(metrics["places"].items(), key=lambda x: x[1]["amount"], reverse=True):
        share = (data["amount"] / tot_rev) * 100
        print(f"{pl:<15} | {format_currency(data['amount']):<16} | {data['units']:<8} | {data['txns']:<7} | {share:>6.2f}%")

    # 4. Yearly Trends
    print("\n[+] YEAR-OVER-YEAR SALES TRENDS")
    print(f"{'Year':<15} | {'Total Sales':<16} | {'Units':<8} | {'Orders':<7} | {'Share %':<8}")
    print("-" * 65)
    for yr in sorted(metrics["years"].keys()):
        data = metrics["years"][yr]
        share = (data["amount"] / tot_rev) * 100
        print(f"{yr:<15} | {format_currency(data['amount']):<16} | {data['units']:<8} | {data['txns']:<7} | {share:>6.2f}%")

    print("\n" + "=" * 72)
    print("Key Takeaway:")
    top_rep = max(metrics["sales_persons"].items(), key=lambda x: x[1]["amount"])
    top_prod = max(metrics["products"].items(), key=lambda x: x[1]["amount"])
    top_zone = max(metrics["places"].items(), key=lambda x: x[1]["amount"])
    print(f"  * Top Salesperson : {top_rep[0]} ({format_currency(top_rep[1]['amount'])}, {(top_rep[1]['amount']/tot_rev)*100:.1f}%)")
    print(f"  * Top Product     : {top_prod[0]} ({format_currency(top_prod[1]['amount'])}, {(top_prod[1]['amount']/tot_rev)*100:.1f}%)")
    print(f"  * Top Market      : {top_zone[0]} ({format_currency(top_zone[1]['amount'])}, {(top_zone[1]['amount']/tot_rev)*100:.1f}%)")
    print("=" * 72)


def main():
    records = get_dataset()
    issues = validate_dataset(records)
    metrics = compute_metrics(records)
    print_dashboard(metrics, issues)


if __name__ == "__main__":
    main()
